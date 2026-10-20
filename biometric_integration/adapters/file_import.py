"""File import - the fallback that works with every machine ever made.

* eSSL / ZKTeco USB download:  1_attlog.dat / GLG_001.TXT (tab separated)
* CSV export from eTimeTrackLite, BioStar, COSEC, Realtime, any software
* Excel (.xlsx) export

Columns are detected automatically (same aliases as the webhook) or can be
named explicitly on the Biometric Log Import form. Separate Date and Time
columns are also supported.
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime, time

from biometric_integration.adapters import adms
from biometric_integration.adapters.base import RawPunch, clean_user_id, parse_datetime, punch_type_from_value
from biometric_integration.adapters.generic import DEVICE_ALIASES, TIME_ALIASES, TYPE_ALIASES, USER_ALIASES

DATE_ALIASES = ["date", "punch_date", "att_date", "log_date", "attendance_date"]
TIME_ONLY_ALIASES = ["time", "punch_time", "in_time", "att_time", "log_time"]


def _norm(text) -> str:
	return str(text or "").strip().lower().replace(" ", "_").replace("-", "_").replace(".", "")


def _find_col(headers: list[str], explicit: str | None, aliases: list[str]) -> int | None:
	normed = [_norm(h) for h in headers]
	if explicit:
		if _norm(explicit) in normed:
			return normed.index(_norm(explicit))
		return None
	for alias in aliases:
		if _norm(alias) in normed:
			return normed.index(_norm(alias))
	return None


def _rows_from_bytes(content: bytes, filename: str) -> list[list]:
	name = filename.lower()
	if name.endswith((".xlsx", ".xlsm")):
		from openpyxl import load_workbook

		wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
		ws = wb.active
		return [list(r) for r in ws.iter_rows(values_only=True)]
	if name.endswith(".xls"):
		import xlrd

		book = xlrd.open_workbook(file_contents=content)
		sheet = book.sheet_by_index(0)
		rows = []
		for r in range(sheet.nrows):
			row = []
			for c in range(sheet.ncols):
				cell = sheet.cell(r, c)
				if cell.ctype == 3:  # date
					row.append(datetime(*xlrd.xldate_as_tuple(cell.value, book.datemode)))
				else:
					row.append(cell.value)
			rows.append(row)
		return rows
	text = content.decode("utf-8-sig", "ignore")
	sample = text[:4096]
	try:
		dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
	except csv.Error:
		dialect = csv.excel
	return [row for row in csv.reader(io.StringIO(text), dialect)]


def _combine(date_value, time_value, date_format):
	if isinstance(date_value, datetime) and time_value in (None, ""):
		return date_value
	if isinstance(date_value, datetime):
		date_value = date_value.date()
	if isinstance(time_value, datetime):
		time_value = time_value.time()
	if isinstance(date_value, date) and isinstance(time_value, time):
		return datetime.combine(date_value, time_value)
	d = date_value.isoformat() if isinstance(date_value, date) else str(date_value).strip()
	t = time_value.strftime("%H:%M:%S") if isinstance(time_value, time) else str(time_value or "").strip()
	return parse_datetime(f"{d} {t}".strip(), date_format=date_format)


def parse_file(
	content: bytes,
	filename: str,
	server_tz=None,
	user_col: str | None = None,
	time_col: str | None = None,
	date_col: str | None = None,
	type_col: str | None = None,
	device_col: str | None = None,
	date_format: str | None = None,
):
	"""Returns (list of (device_sn_or_None, RawPunch), list of error strings)."""
	name = filename.lower()
	text_head = content[:2000].decode("utf-8", "ignore")

	# eSSL / ZKTeco USB attlog: no header, tab separated, PIN first
	if name.endswith((".dat", ".txt")) or (
		"\t" in text_head and adms._LINE_RE.match(text_head.splitlines()[0] if text_head else "")
	):
		punches, bad = adms.parse_attlog(content.decode("utf-8", "ignore"), server_tz)
		if punches:
			return [(None, p) for p in punches], [f"Could not read line: {b}" for b in bad[:50]]

	rows = [r for r in _rows_from_bytes(content, filename) if any(str(c or "").strip() for c in r)]
	if not rows:
		return [], ["File is empty"]

	# find header row within first 10 rows (some exports have a title on top)
	header_idx, cols = None, {}
	for i, row in enumerate(rows[:10]):
		headers = [str(c or "") for c in row]
		u = _find_col(headers, user_col, USER_ALIASES + ["employee_id", "emp_id", "empid", "id", "code", "user_no", "ac_no", "acno"])
		t = _find_col(headers, time_col, TIME_ALIASES)
		d = _find_col(headers, date_col, DATE_ALIASES)
		if u is not None and (t is not None or d is not None):
			header_idx = i
			cols = {
				"user": u,
				"time": t,
				"date": d,
				"type": _find_col(headers, type_col, TYPE_ALIASES),
				"device": _find_col(headers, device_col, DEVICE_ALIASES),
			}
			if d is not None and t is not None and t == d:
				cols["date"] = None
			break
	if header_idx is None:
		return [], [
			"Could not find the User ID and Time columns. Put the column names in the "
			"'Column Mapping' section and try again. First row seen: " + ", ".join(str(c) for c in rows[0])
		]

	out, errors = [], []
	for n, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):

		def cell(key):
			idx = cols.get(key)
			return row[idx] if idx is not None and idx < len(row) else None

		user = clean_user_id(cell("user"))
		if not user:
			continue
		try:
			time_value = cell("time")
			full_dt = isinstance(time_value, datetime) and time_value.year > 1900
			if not full_dt and isinstance(time_value, str) and any(s in time_value for s in ("-", "/", ".")) and len(time_value.strip()) > 10:
				full_dt = True
			if cols.get("date") is not None and not full_dt:
				punch_time = _combine(cell("date"), time_value, date_format)
			else:
				value = cell("time")
				punch_time = value if isinstance(value, datetime) else parse_datetime(value, server_tz, date_format)
			punch_time = punch_time.replace(microsecond=0, tzinfo=None)
		except (ValueError, TypeError) as e:
			errors.append(f"Row {n}: {e}")
			continue
		type_value = cell("type")
		device = cell("device")
		out.append(
			(
				str(device).strip() if device not in (None, "") else None,
				RawPunch(
					user_id=user,
					punch_time=punch_time,
					punch_type=punch_type_from_value(type_value),
					status_code="" if type_value is None else str(type_value),
					raw=", ".join("" if c is None else str(c) for c in row)[:1000],
				),
			)
		)
	return out, errors
